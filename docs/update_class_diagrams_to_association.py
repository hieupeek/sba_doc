#!/usr/bin/env python3
"""
update_class_diagrams_to_association.py

Updates all Backend Class Diagrams:
- Replaces incorrect Aggregation (o--) between Service -> Repository and Entity -> Entity
  with proper UML Association (-->).
- Regenerates all 5 class diagram PNGs using PlantUML via Kroki.
- Replaces the 5 embedded images in HMS_SoftwareDesignTemplate-6.docx.
- Updates the explanatory text paragraphs in the Word document.
"""

import os
import urllib.request
import docx
import xml.etree.ElementTree as ET

DOCX_PATH = "/Users/hieu/Desktop/sba/docs/HMS_SoftwareDesignTemplate-6.docx"
MD_PATH = "/Users/hieu/Desktop/sba/docs/Class + Sequence/Component Structure Design.md"
OUTPUT_DIR = "/Users/hieu/Desktop/sba/docs/Class + Sequence"

DIAGRAMS = {
    "class_backend_42_auth.png": {
        "title": "4.2: Account & Authentication Management",
        "puml": """@startuml
skinparam shadowed false
hide circle

' 1. Declaration of Classes
class UserController <<controller>>
class AuthenticationController <<controller>>
class JwtAuthenticationFilter <<security>>
class JwtTokenProvider

interface UserService <<interface>>
interface AuthenticationService <<interface>>

class UserServiceImpl <<service>>
class AuthenticationServiceImpl <<service>>

class UserRepository <<repository>>
class RoleRepository <<repository>>

class User <<entity>>
class Role <<entity>>

' 2. Realization (..|>)
UserServiceImpl ..|> UserService
AuthenticationServiceImpl ..|> AuthenticationService

' 3. Dependency (..>)
UserController ..> UserService
AuthenticationController ..> AuthenticationService
JwtAuthenticationFilter ..> JwtTokenProvider
JwtAuthenticationFilter ..> AuthenticationService
UserRepository ..> User
RoleRepository ..> Role

' 4. Association (-->)
UserServiceImpl --> UserRepository
AuthenticationServiceImpl --> UserRepository
AuthenticationServiceImpl --> RoleRepository
User --> Role

@enduml"""
    },

    "class_backend_43_room.png": {
        "title": "4.3: Room & Facility Management",
        "puml": """@startuml
skinparam shadowed false
hide circle

' 1. Declaration of Classes
class RoomController <<controller>>
class HousekeepingController <<controller>>
class RoomMaintenanceController <<controller>>

interface RoomService <<interface>>
interface HousekeepingService <<interface>>
interface RoomMaintenanceService <<interface>>

class RoomServiceImpl <<service>>
class HousekeepingServiceImpl <<service>>
class RoomMaintenanceServiceImpl <<service>>

class RoomRepository <<repository>>
class RoomTypeRepository <<repository>>
class RoomMaintenanceRepository <<repository>>

class Room <<entity>>
class Room_Type <<entity>>
class Room_Maintenance <<entity>>
class User <<entity>>

' 2. Realization (..|>)
RoomServiceImpl ..|> RoomService
HousekeepingServiceImpl ..|> HousekeepingService
RoomMaintenanceServiceImpl ..|> RoomMaintenanceService

' 3. Dependency (..>)
RoomController ..> RoomService
HousekeepingController ..> HousekeepingService
RoomMaintenanceController ..> RoomMaintenanceService
RoomRepository ..> Room
RoomTypeRepository ..> Room_Type
RoomMaintenanceRepository ..> Room_Maintenance

' 4. Association (-->)
RoomServiceImpl --> RoomRepository
RoomServiceImpl --> RoomTypeRepository
HousekeepingServiceImpl --> RoomRepository
RoomMaintenanceServiceImpl --> RoomMaintenanceRepository
RoomMaintenanceServiceImpl --> RoomRepository
Room --> Room_Type
Room_Maintenance --> Room
Room_Maintenance --> User

@enduml"""
    },

    "class_backend_44_reservation.png": {
        "title": "4.4: Reservation & Front-Desk Operations",
        "puml": """@startuml
skinparam shadowed false
hide circle

' 1. Declaration of Classes
class ReservationController <<controller>>
class FrontDeskController <<controller>>

interface ReservationService <<interface>>
interface FrontDeskService <<interface>>

class ReservationServiceImpl <<service>>
class FrontDeskServiceImpl <<service>>

class ReservationRepository <<repository>>
class ReservationDetailRepository <<repository>>
class GuestStayRepository <<repository>>
class StayRoomAssignmentRepository <<repository>>

class Reservation <<entity>>
class Reservation_Detail <<entity>>
class Guest_Stay <<entity>>
class Stay_Room_Assignment <<entity>>
class Customer <<entity>>
class Room <<entity>>
class Room_Type <<entity>>

' 2. Realization (..|>)
ReservationServiceImpl ..|> ReservationService
FrontDeskServiceImpl ..|> FrontDeskService

' 3. Dependency (..>)
ReservationController ..> ReservationService
FrontDeskController ..> FrontDeskService
ReservationRepository ..> Reservation
ReservationDetailRepository ..> Reservation_Detail
GuestStayRepository ..> Guest_Stay
StayRoomAssignmentRepository ..> Stay_Room_Assignment

' 4. Association (-->)
ReservationServiceImpl --> ReservationRepository
ReservationServiceImpl --> ReservationDetailRepository
FrontDeskServiceImpl --> GuestStayRepository
FrontDeskServiceImpl --> StayRoomAssignmentRepository
FrontDeskServiceImpl --> ReservationRepository
Reservation --> Customer
Guest_Stay --> Reservation
Reservation_Detail --> Room_Type
Stay_Room_Assignment --> Room

' 5. Composition (*--)
Reservation *-- Reservation_Detail
Guest_Stay *-- Stay_Room_Assignment

@enduml"""
    },

    "class_backend_45_billing.png": {
        "title": "4.5: Ancillary Services & Billing Management",
        "puml": """@startuml
skinparam shadowed false
hide circle

' 1. Declaration of Classes
class HotelServiceController <<controller>>
class ServiceOrderController <<controller>>
class BillingController <<controller>>
class PaymentController <<controller>>

interface HotelServiceService <<interface>>
interface ServiceOrderService <<interface>>
interface BillingService <<interface>>
interface PaymentService <<interface>>

class HotelServiceServiceImpl <<service>>
class ServiceOrderServiceImpl <<service>>
class BillingServiceImpl <<service>>
class PaymentServiceImpl <<service>>

class HotelServiceRepository <<repository>>
class ServiceOrderRepository <<repository>>
class FolioRepository <<repository>>
class AdditionalChargeRepository <<repository>>
class InvoiceRepository <<repository>>
class PaymentRepository <<repository>>
class PaymentAttemptRepository <<repository>>
class VoucherRepository <<repository>>

class Hotel_Service <<entity>>
class Service_Order <<entity>>
class Folio <<entity>>
class Additional_Charge <<entity>>
class Invoice <<entity>>
class Payment <<entity>>
class Payment_Attempt <<entity>>
class Voucher <<entity>>
class Guest_Stay <<entity>>
class Reservation <<entity>>
class User <<entity>>

' 2. Realization (..|>)
HotelServiceServiceImpl ..|> HotelServiceService
ServiceOrderServiceImpl ..|> ServiceOrderService
BillingServiceImpl ..|> BillingService
PaymentServiceImpl ..|> PaymentService

' 3. Dependency (..>)
HotelServiceController ..> HotelServiceService
ServiceOrderController ..> ServiceOrderService
BillingController ..> BillingService
PaymentController ..> PaymentService
HotelServiceRepository ..> Hotel_Service
ServiceOrderRepository ..> Service_Order
FolioRepository ..> Folio
AdditionalChargeRepository ..> Additional_Charge
InvoiceRepository ..> Invoice
PaymentRepository ..> Payment
PaymentAttemptRepository ..> Payment_Attempt
VoucherRepository ..> Voucher

' 4. Association (-->)
HotelServiceServiceImpl --> HotelServiceRepository
ServiceOrderServiceImpl --> ServiceOrderRepository
ServiceOrderServiceImpl --> AdditionalChargeRepository
BillingServiceImpl --> FolioRepository
BillingServiceImpl --> AdditionalChargeRepository
BillingServiceImpl --> InvoiceRepository
BillingServiceImpl --> VoucherRepository
PaymentServiceImpl --> PaymentRepository
PaymentServiceImpl --> InvoiceRepository
PaymentServiceImpl --> PaymentAttemptRepository
Service_Order --> Hotel_Service
Service_Order --> Guest_Stay
Folio --> Guest_Stay
Additional_Charge --> User
Invoice --> Folio
Invoice --> Voucher
Payment --> Invoice
Payment --> Reservation

' 5. Composition (*--)
Folio *-- Additional_Charge
Payment *-- Payment_Attempt

@enduml"""
    },

    "class_backend_46_reporting.png": {
        "title": "4.6: Reviews, Audit & Reporting Management",
        "puml": """@startuml
skinparam shadowed false
hide circle

' 1. Declaration of Classes
class ReviewController <<controller>>
class ReportController <<controller>>

interface ReviewService <<interface>>
interface ReportService <<interface>>

class ReviewServiceImpl <<service>>
class ReportServiceImpl <<service>>

class ReviewRepository <<repository>>
class ReservationRepository <<repository>>
class RoomRepository <<repository>>

class Review <<entity>>
class Customer <<entity>>
class Guest_Stay <<entity>>
class Reservation <<entity>>
class Room <<entity>>

' 2. Realization (..|>)
ReviewServiceImpl ..|> ReviewService
ReportServiceImpl ..|> ReportService

' 3. Dependency (..>)
ReviewController ..> ReviewService
ReportController ..> ReportService
ReviewRepository ..> Review
ReservationRepository ..> Reservation
RoomRepository ..> Room
ReportServiceImpl ..> RoomRepository

' 4. Association (-->)
ReviewServiceImpl --> ReviewRepository
ReportServiceImpl --> ReservationRepository
ReportServiceImpl --> RoomRepository
Review --> Customer
Review --> Guest_Stay

@enduml"""
    }
}

def render_plantuml_kroki(puml_text):
    req = urllib.request.Request(
        "https://kroki.io/plantuml/png",
        data=puml_text.encode("utf-8"),
        headers={
            "Content-Type": "text/plain; charset=utf-8",
            "User-Agent": "Mozilla/5.0"
        }
    )
    with urllib.request.urlopen(req, timeout=15) as resp:
        return resp.read()

def main():
    print("=== Step 1: Rendering updated PlantUML diagrams via Kroki ===")
    rendered_images = {}
    for filename, diag_info in DIAGRAMS.items():
        print(f"Rendering {filename} ({diag_info['title']})...")
        png_bytes = render_plantuml_kroki(diag_info["puml"])
        out_path = os.path.join(OUTPUT_DIR, filename)
        with open(out_path, "wb") as f:
            f.write(png_bytes)
        rendered_images[filename] = png_bytes
        print(f"  -> Saved {out_path} ({len(png_bytes)} bytes)")

    print("\n=== Step 2: Updating Component Structure Design.md ===")
    new_md_content = """# Backend Class/Component Structure Design (PlantUML)

Tài liệu thiết kế cấu trúc thành phần tầng Backend cho toàn bộ 5 phân hệ của hệ thống Hotel Management System (HMS).
Được thiết kế chuẩn theo 6 loại quan hệ UML trong SKILL.md và đồng nhất 100% với CSDL quan hệ trong Section 5:
- **Inheritance (`--|>`):** Kế thừa lớp cha / repository base.
- **Realization (`..|>`):** ServiceImpl hiện thực hóa Service interface.
- **Dependency (`..>`):** Controller sử dụng Service, Repository trả về Entity.
- **Association (`-->`):** Liên kết có hướng chuẩn mực (Service giữ tham chiếu điều hướng Repository thông qua Dependency Injection; Entity tham chiếu Entity khác thông qua khóa ngoại, ví dụ `User --> Role`, `Room --> Room_Type`, `Review --> Guest_Stay`, `Invoice --> Voucher`, `Payment --> Reservation`).
- **Composition (`*--`):** Has-a chặt chẽ phụ thuộc vòng đời (Reservation sở hữu Reservation_Detail, Guest_Stay sở hữu Stay_Room_Assignment, Folio sở hữu Additional_Charge, Payment sở hữu Payment_Attempt).
- **Aggregation (`o--`):** Has-a sở hữu lỏng lẻo toàn thể - bộ phận (Whole-Part độc lập).

---

## 1. Mục 4.2: Account & Authentication Management

```plantuml
""" + DIAGRAMS["class_backend_42_auth.png"]["puml"] + """
```

---

## 2. Mục 4.3: Room & Facility Management (Housekeeping & Maintenance)

```plantuml
""" + DIAGRAMS["class_backend_43_room.png"]["puml"] + """
```

---

## 3. Mục 4.4: Reservation & Front-Desk Operations

```plantuml
""" + DIAGRAMS["class_backend_44_reservation.png"]["puml"] + """
```

---

## 4. Mục 4.5: Ancillary Services & Billing Management

```plantuml
""" + DIAGRAMS["class_backend_45_billing.png"]["puml"] + """
```

---

## 5. Mục 4.6: Reviews, Audit & Reporting Management

```plantuml
""" + DIAGRAMS["class_backend_46_reporting.png"]["puml"] + """
```
"""
    with open(MD_PATH, "w", encoding="utf-8") as f:
        f.write(new_md_content)
    print(f"Updated {MD_PATH}")

    print("\n=== Step 3: Updating Word Document (HMS_SoftwareDesignTemplate-6.docx) ===")
    doc = docx.Document(DOCX_PATH)

    # 3.1: Update image blobs
    # Map paragraph indices to rendered filename
    p_to_file = {
        213: "class_backend_42_auth.png",
        245: "class_backend_43_room.png",
        281: "class_backend_44_reservation.png",
        313: "class_backend_45_billing.png",
        341: "class_backend_46_reporting.png"
    }

    for p_idx, filename in p_to_file.items():
        p = doc.paragraphs[p_idx]
        found_embed = False
        for r in p.runs:
            if "blip" in r._r.xml:
                tree = ET.fromstring(r._r.xml)
                for elem in tree.iter():
                    if elem.tag.endswith("blip"):
                        embed_id = elem.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                        part = doc.part.related_parts[embed_id]
                        # Replace part's blob
                        part._blob = rendered_images[filename]
                        print(f"  -> Replaced image at P{p_idx} ({part.partname}) with {filename} ({len(rendered_images[filename])} bytes)")
                        found_embed = True
                        break
                if found_embed:
                    break

    # 3.2: Update descriptive text paragraphs
    text_updates = {
        212: "The backend architecture strictly adheres to standard Java Spring Boot 3-Tier Enterprise patterns. The diagram uses flat layout and strictly applies standard UML relationship arrows: Realization (..|>) between ServiceImpl and Interface, Dependency (..>) from Controller to Service and Repository to Entity, and Association (-->) for Spring bean repository references and domain entity associations.",
        244: "Figure 4-6 presents the backend class diagram for Room & Facility Management using flat layout and rigorous UML relationships: Realization (..|>) for services, Dependency (..>) for controller and repository interactions, and Association (-->) for repository references and entity associations.",
        280: "Figure 4-11 illustrates the backend class structure for Reservation and Front-Desk operations. Crucially, it utilizes UML Composition (*--) for lifecycle-dependent parent-child entities: Reservation *-- Reservation_Detail and Guest_Stay *-- Stay_Room_Assignment, alongside Realization (..|>) for services and Association (-->) for repository references and domain entity associations.",
        312: "Figure 4-15 illustrates the backend class structure for Ancillary Services and Billing Management. It enforces Composition (*--) for Folio *-- Additional_Charge, Realization (..|>) for all billing and order services, Dependency (..>) for controller endpoints and JPA queries, and Association (-->) for repository references and billing entity associations.",
        340: "Figure 4-18 illustrates the backend class structure for Reviews & Reporting. It uses flat layout with standard UML relationship arrows: Realization (..|>) for reporting services, Dependency (..>) for controller endpoints and query invocations, and Association (-->) for repository references and review entity associations."
    }

    for p_idx, new_text in text_updates.items():
        p = doc.paragraphs[p_idx]
        p.text = new_text
        print(f"  -> Updated text at P{p_idx}")

    doc.save(DOCX_PATH)
    print(f"Saved cleanly: {DOCX_PATH}")

if __name__ == "__main__":
    main()
